"""Check installed hooks, real macOS isolation, and optional installed Pi runtime.

Run from an installed environment: python examples/check_connections.py
Optional: --pi-package /path/to/@earendil-works/pi-coding-agent --node /path/to/node
No model calls, global Agent configuration, or real user documents.
"""

import argparse
import json
import os
import shlex
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook

from filewise.connections import AGENTS, Connections
from filewise.engine import Engine, FilewiseError
from filewise.middleware import Middleware, WatchConfig
from filewise.models import Actor
from filewise.projects import Projects


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pi-package", type=Path)
    parser.add_argument("--node", default="node")
    args = parser.parse_args()
    if sys.platform != "darwin":
        raise SystemExit("Requires macOS sandbox-exec")
    env = dict(os.environ)
    if env.get("PYTHONPATH"):
        env["PYTHONPATH"] = os.pathsep.join(
            str(Path(p).resolve()) for p in env["PYTHONPATH"].split(os.pathsep)
        )
    with tempfile.TemporaryDirectory(prefix="filewise-connections-check-") as directory:
        root = Path(directory).resolve() / "关注文件"
        root.mkdir()
        (root / "notes.md").write_text("100")
        book = Workbook()
        book.active.append(["Pressure", 100])
        book.save(root / "inspection.xlsx")
        book.close()
        owner = Actor(id="owner", roles={"reader", "editor", "reviewer", "publisher"})
        engine = Engine(root / ".filewise/filewise.db")
        projects = Projects(engine)
        projects.create({"id": "files", "name": "Synthetic files"}, owner, root)
        middleware = Middleware(projects)
        connections = Connections(middleware)
        middleware.configure("files", WatchConfig(), owner)
        report = {}
        for agent in AGENTS:
            connections.install("files", agent, owner)
            assert connections.self_test("files", agent, owner)["passed"]
            status = next(a for a in connections.status("files", owner)["agents"] if a["agent"] == agent)
            assert status["state"] == "waiting", "Self-test must not count as an observed native event"
        report["three_connection_self_tests"] = "PASS; waiting for native events"

        for agent, pressure in (("codex", 120), ("claude", 130)):
            config = json.loads((root / AGENTS[agent][1]).read_text())

            def event(phase, inputs=None):
                command = config["hooks"][phase][0]["hooks"][0]["command"]
                result = subprocess.run(
                    ["/bin/sh", "-c", command],
                    cwd=root,
                    env=env,
                    text=True,
                    capture_output=True,
                    input=json.dumps(
                        {
                            "hook_event_name": phase,
                            "session_id": agent + "-fixture",
                            "cwd": str(root),
                            "tool_name": "Bash",
                            "tool_input": inputs or {},
                        }
                    ),
                    timeout=30,
                )
                assert result.returncode == 0, result.stderr
                return json.loads(result.stdout)["hookSpecificOutput"]

            event("SessionStart")
            before = (root / "inspection.xlsx").read_bytes()
            # These are native protocol fixtures, not a live Codex/Claude model session.
            script = "from pathlib import Path\nfrom openpyxl import load_workbook\n"
            for path in (root / "notes.md", engine.path):
                script += f"try: Path({str(path)!r}).read_bytes()\nexcept PermissionError: pass\nelse: raise AssertionError('raw read allowed')\n"
                script += f"try: Path({str(path)!r}).open('ab')\nexcept PermissionError: pass\nelse: raise AssertionError('raw write allowed')\n"
            script += f"b=load_workbook('inspection.xlsx'); b.active['B1']={pressure}; b.save('inspection.xlsx'); b.close()\n"
            original_command = shlex.join([sys.executable, "-c", script])
            result = event("PreToolUse", {"command": original_command})
            assert "updatedInput" in result, result
            wrapped = result["updatedInput"]["command"]
            job = Path(shlex.split(wrapped)[4])
            assert not (job / "start").exists()
            assert (job / "done").read_bytes() == b"", "Command ran before the tool started"
            # Tool processes must not need to write the private DB; the after-tool hook captures changes.
            outer = (
                "(version 1)(allow default)(deny file-write* (literal " + json.dumps(str(engine.path)) + "))"
            )
            execution = subprocess.run(
                ["sandbox-exec", "-p", outer, "/bin/sh", "-c", wrapped],
                cwd=root,
                env=env,
                capture_output=True,
                text=True,
                timeout=30,
            )
            assert execution.returncode == 0, execution.stderr
            event("PostToolUseFailure" if agent == "claude" else "PostToolUse")
            assert (root / "inspection.xlsx").read_bytes() == before
            rid = middleware.config("files", owner)["last_event"]["release_id"]
            active = projects.detail("files", owner)["active_release"]
            try:
                middleware.apply("files", rid, active, owner)
            except FilewiseError:
                pass
            else:
                raise AssertionError("Unapproved writeback succeeded")
            projects.approve("files", rid, owner)
            middleware.apply("files", rid, active, owner)
            book = load_workbook(root / "inspection.xlsx")
            assert book.active["B1"].value == pressure
            book.close()
            report[agent + "_protocol_and_excel_writeback"] = "PASS"

        waiting = event("PreToolUse", {"command": "sleep 20"})["updatedInput"]["command"]
        job = Path(shlex.split(waiting)[4])
        tool = subprocess.Popen(
            shlex.split(waiting), cwd=root, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
        )
        for _ in range(100):
            if (job / "start").exists():
                break
            time.sleep(0.05)
        else:
            raise AssertionError("Native tool did not start")
        tool.terminate()
        tool.wait(timeout=5)
        for _ in range(120):
            if (job / "done").read_bytes():
                break
            time.sleep(0.05)
        else:
            raise AssertionError("Worker did not stop after native tool cancellation")
        assert int((job / "done").read_text()) != 0
        report["native_tool_cancellation"] = "PASS"

        if args.pi_package:
            # Use Pi's actual TS loader, event runner, and write tool; no model provider needed.
            module = (args.pi_package.resolve() / "dist/core").as_uri()
            script = f"""
import assert from 'node:assert/strict';
import {{readFileSync}} from 'node:fs';
import {{loadExtensions}} from '{module}/extensions/loader.js';
import {{ExtensionRunner}} from '{module}/extensions/runner.js';
import {{createWriteTool}} from '{module}/tools/write.js';
const root = {json.dumps(str(root))};
const loaded = await loadExtensions([root + '/.pi/extensions/filewise.ts'], root);
assert.deepEqual(loaded.errors, []);
const runner = new ExtensionRunner(loaded.extensions, loaded.runtime, root, {{getSessionId: () => 'pi-runtime-check'}}, {{}});
const errors = []; runner.onError(e => errors.push(e));
await runner.emit({{type: 'session_start'}});
const input = {{path: 'notes.md', content: 'Pi native tool, pending review'}};
const event = {{type: 'tool_call', toolName: 'write', toolCallId: 'test-write', input}};
const decision = await runner.emitToolCall(event);
assert.ok(!decision?.block, JSON.stringify(decision));
assert.ok(input.path.startsWith(root + '/.filewise/workspaces/'));
const result = await createWriteTool(root).execute('test-write', input);
const post = await runner.emitToolResult({{...event, type: 'tool_result', ...result, isError: false}});
assert.ok(post.content.some(c => c.text?.includes('Filewise')));
assert.equal(readFileSync(root + '/notes.md', 'utf8'), '100');
assert.deepEqual(errors, []);
console.log('Pi loader + event runner + native write tool: PASS');
"""
            result = subprocess.run(
                [args.node, "--input-type=module", "-e", script],
                env=env,
                cwd=directory,
                text=True,
                capture_output=True,
                timeout=45,
            )
            assert result.returncode == 0, result.stderr
            report["pi_installed_runtime"] = "PASS"
        else:
            report["pi_installed_runtime"] = "NOT RUN; pass --pi-package to verify"
        for agent in AGENTS:
            connections.disconnect("files", agent, owner)
            assert not (root / AGENTS[agent][1]).exists()
        report["disconnect_configuration_restored"] = "PASS"
        middleware.close()
        print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
