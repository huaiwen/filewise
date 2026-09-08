"""Deterministic extraction. Fragments are evidence, never approved business IR."""

import csv
import io
import zipfile
from pathlib import Path

from .engine import FilewiseError

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 2_000_000
MAX_FRAGMENTS = 20_000
TEXT_TYPES = {".txt", ".md", ".csv"}
OFFICE_TYPES = {".docx", ".xlsx", ".pptx"}


def extract(name: str, body: bytes) -> tuple[str, list[dict]]:
    """Bound compressed input/output; preserve stable coordinates and exact text."""
    suffix = Path(name).suffix.lower()
    if not body or len(body) > MAX_BYTES:
        raise FilewiseError("Source must be nonempty and at most 10 MiB", 413)
    if suffix not in TEXT_TYPES | OFFICE_TYPES | {".pdf"}:
        raise FilewiseError("Supported formats: txt, md, csv; documents extra: pdf, docx, xlsx, pptx", 415)
    fragments, text_size = [], 0

    def add(locator, text):
        nonlocal text_size
        if not text.strip():
            return
        text_size += len(text)
        if text_size > MAX_TEXT or len(fragments) >= MAX_FRAGMENTS:
            raise FilewiseError("Extracted content exceeds limit", 413)
        fragments.append({"locator": locator, "text": text})

    try:
        if suffix in OFFICE_TYPES:
            with zipfile.ZipFile(io.BytesIO(body)) as archive:
                entries = archive.infolist()
                if len(entries) > 10_000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                    raise FilewiseError("Expanded Office archive exceeds limit", 413)
        if suffix in TEXT_TYPES:
            text = body.decode("utf-8-sig")
            if "\x00" in text:
                raise FilewiseError("Binary content is not UTF-8 text", 415)
            if suffix == ".csv":
                for row, cells in enumerate(csv.reader(io.StringIO(text, newline=""), strict=True), 1):
                    for col, cell in enumerate(cells, 1):
                        add(f"row:{row}/col:{col}", cell)
            else:
                for line, content in enumerate(text.splitlines(), 1):
                    add(f"line:{line}", content)
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(body))
            if reader.is_encrypted or len(reader.pages) > 500:
                raise FilewiseError("Encrypted PDFs or PDFs over 500 pages are unsupported", 415)
            for page, content in enumerate(reader.pages, 1):
                add(f"page:{page}", content.extract_text() or "")
        elif suffix == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(body))
            for index, paragraph in enumerate(doc.paragraphs, 1):
                add(f"paragraph:{index}", paragraph.text)
            for table, content in enumerate(doc.tables, 1):
                for row, cells in enumerate(content.rows, 1):
                    for col, cell in enumerate(cells.cells, 1):
                        add(f"table:{table}/row:{row}/col:{col}", cell.text)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
            try:
                for sheet_no, sheet in enumerate(book.worksheets, 1):
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > 200_000:
                        raise FilewiseError("Worksheet exceeds 200,000 cells", 413)
                    for row in sheet:
                        for cell in row:
                            if cell.value is not None:
                                add(f"sheet:{sheet_no}/cell:{cell.coordinate}", str(cell.value))
            finally:
                book.close()
        else:
            from pptx import Presentation

            for slide, content in enumerate(Presentation(io.BytesIO(body)).slides, 1):
                for shape, item in enumerate(content.shapes, 1):
                    if item.has_text_frame:
                        add(f"slide:{slide}/shape:{shape}", item.text)
                    if item.has_table:
                        for row, cells in enumerate(item.table.rows, 1):
                            for col, cell in enumerate(cells.cells, 1):
                                add(f"slide:{slide}/shape:{shape}/row:{row}/col:{col}", cell.text)
    except ImportError as exc:
        raise FilewiseError("Install filewise-engine[documents] to read this format", 415) from exc
    except FilewiseError:
        raise
    except Exception as exc:
        raise FilewiseError("Cannot parse source; check its format and encoding", 422) from exc
    if not fragments:
        raise FilewiseError("No extractable text; OCR/vision is not implemented", 422)
    return f"filewise/{suffix[1:]}-v1", fragments


def ingest(engine, scope_id, name, body, actor, acl=None):
    # Never interpret a supplied filename as a filesystem destination.
    name = name.replace("\\", "/").rsplit("/", 1)[-1]
    if not name or len(name) > 255:
        raise FilewiseError("Filename must contain 1–255 characters")
    parser, fragments = extract(name, body)
    return engine.add_source(scope_id, name, body, fragments, parser, actor, acl)
