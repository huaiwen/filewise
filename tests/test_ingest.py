import importlib.util
import io
import unittest
import zipfile

from filewise import FilewiseError
from filewise.ingest import MAX_BYTES, extract


class IngestTests(unittest.TestCase):
    def test_text_and_multiline_csv_coordinates(self):
        parser, fragments = extract("sample.md", "\ufeff标题\r\n\r\n原文证据\n".encode())
        self.assertEqual(parser, "filewise/md-v1")
        self.assertEqual(
            fragments, [{"locator": "line:1", "text": "标题"}, {"locator": "line:3", "text": "原文证据"}]
        )
        _, cells = extract("sample.csv", b'name,value\n"multi\nline","a,b"\n')
        self.assertEqual(
            cells[-2:],
            [{"locator": "row:2/col:1", "text": "multi\nline"}, {"locator": "row:2/col:2", "text": "a,b"}],
        )

    def test_reject_empty_binary_large_and_malformed(self):
        for name, body in [
            ("a.txt", b""),
            ("a.txt", b"x" * (MAX_BYTES + 1)),
            ("a.exe", b"text"),
            ("a.txt", b"\xff"),
            ("a.txt", b"\x00"),
            ("a.txt", b" \n"),
            ("a.csv", b'"unclosed'),
            ("a.docx", b"not a zip"),
            ("a.png", b"image"),
        ]:
            with self.subTest(name=name, size=len(body)), self.assertRaises(FilewiseError):
                extract(name, body)

    def test_fragment_and_expanded_archive_limits(self):
        with self.assertRaises(FilewiseError):
            extract("large.txt", b"a\n" * 20_001)
        with self.assertRaises(FilewiseError):
            extract("large.txt", b"a" * 2_000_001)
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", b"0" * (51 * 1024 * 1024))
        with self.assertRaises(FilewiseError) as caught:
            extract("large.docx", buffer.getvalue())
        self.assertEqual(caught.exception.status, 413)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "documents extra not installed")
    def test_docx_paragraph_and_table(self):
        from docx import Document

        doc, buffer = Document(), io.BytesIO()
        doc.add_paragraph("Synthetic paragraph")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic cell"
        doc.save(buffer)
        _, fragments = extract("sample.docx", buffer.getvalue())
        self.assertEqual(
            fragments,
            [
                {"locator": "paragraph:1", "text": "Synthetic paragraph"},
                {"locator": "table:1/row:1/col:1", "text": "Synthetic cell"},
            ],
        )

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "documents extra not installed")
    def test_xlsx_formula_is_not_executed(self):
        from openpyxl import Workbook

        book, buffer = Workbook(), io.BytesIO()
        book.active["A1"] = "Synthetic"
        book.active["B2"] = "=1+2"
        book.save(buffer)
        book.close()
        _, fragments = extract("sample.xlsx", buffer.getvalue())
        self.assertEqual(fragments[-1], {"locator": "sheet:1/cell:B2", "text": "=1+2"})

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "documents extra not installed")
    def test_pptx_text(self):
        from pptx import Presentation
        from pptx.util import Inches

        slides, buffer = Presentation(), io.BytesIO()
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Synthetic slide"
        slides.save(buffer)
        self.assertEqual(
            extract("sample.pptx", buffer.getvalue())[1],
            [{"locator": "slide:1/shape:1", "text": "Synthetic slide"}],
        )

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "documents extra not installed")
    def test_pdf_text_and_scanned_refusal(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer, buffer = PdfWriter(), io.BytesIO()
        page = writer.add_blank_page(width=300, height=300)
        writer.write(buffer)
        with self.assertRaises(FilewiseError) as caught:
            extract("scan.pdf", buffer.getvalue())
        self.assertIn("OCR/vision", str(caught.exception))
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 10 100 Td (Synthetic evidence) Tj ET")
        page[NameObject("/Contents")] = stream
        buffer = io.BytesIO()
        writer.write(buffer)
        fragments = extract("sample.pdf", buffer.getvalue())[1]
        self.assertEqual(fragments[0]["locator"], "page:1")
        self.assertIn("Synthetic evidence", fragments[0]["text"])


if __name__ == "__main__":
    unittest.main()
